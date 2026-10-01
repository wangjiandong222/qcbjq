from __future__ import annotations

import csv
import json
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy.orm import Session

from app.models.entities import ImportBatch, WorkOrder


COLUMN_MAP = {
    "工单编号": "order_no",
    "工单类型": "order_type",
    "问题分类": "problem_category",
    "标签": "tags",
    "标题": "title",
    "主要内容": "content",
    "工单状态": "status",
    "来电人": "caller_name",
    "来电人电话/账号": "caller_phone",
    "被反映区": "district",
    "被反映街乡镇": "town",
    "办理结果": "handling_result",
    "回复内容": "reply_content",
    "处理受理方式": "handling_method",
    "主办单位": "host_unit",
    "来电时间": "received_at",
    "办结时间": "closed_at",
    "企业名称": "company_name",
    "是否解决": "is_resolved",
    "是否满意": "satisfaction",
    "工单性质": "order_nature",
    "村/社区": "community",
    "小区点位": "location_point",
}

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".csv"}

COLUMN_ALIASES = {
    "工单编号": ["工单编号", "工单号", "工单编码", "编号", "诉求编号", "受理编号"],
    "工单类型": ["工单类型", "诉求类型", "事项类型"],
    "问题分类": ["问题分类", "问题类别", "事项分类", "诉求分类", "来电类别"],
    "标签": ["标签", "关键词标签", "工单标签"],
    "标题": ["标题", "工单标题", "诉求标题", "问题标题", "来电标题"],
    "主要内容": ["主要内容", "内容", "诉求内容", "反映内容", "来电内容", "问题描述", "工单内容", "投诉内容", "咨询内容"],
    "工单状态": ["工单状态", "状态", "办理状态"],
    "来电人": ["来电人", "反映人", "诉求人", "联系人", "来话人"],
    "来电人电话/账号": ["来电人电话/账号", "来电人电话", "联系电话", "联系方式", "电话", "账号", "来电号码"],
    "被反映区": ["被反映区", "反映区", "区县", "所在区"],
    "被反映街乡镇": ["被反映街乡镇", "街乡镇", "乡镇街道", "街道乡镇", "被反映街道", "所在街乡镇"],
    "办理结果": ["办理结果", "处理结果", "承办结果"],
    "回复内容": ["回复内容", "答复内容", "办理回复"],
    "处理受理方式": ["处理受理方式", "受理方式", "处理方式"],
    "主办单位": ["主办单位", "承办单位", "办理单位", "处理单位", "责任单位"],
    "来电时间": ["来电时间", "来话时间", "受理时间", "创建时间", "登记时间", "投诉时间", "反映时间", "工单时间"],
    "办结时间": ["办结时间", "完成时间", "结案时间", "处理时间"],
    "企业名称": ["企业名称", "单位名称", "商户名称", "公司名称"],
    "是否解决": ["是否解决", "解决情况", "是否办结解决"],
    "是否满意": ["是否满意", "满意度", "评价结果"],
    "工单性质": ["工单性质", "诉求性质", "事项性质"],
    "村/社区": ["村/社区", "村社区", "社区", "村居"],
    "小区点位": ["小区点位", "点位", "地址", "详细地址", "发生地点", "问题点位", "小区"],
}


@dataclass(frozen=True)
class HeaderInfo:
    name: str
    canonical: str
    standard_key: str | None = None


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    text = _clean(value)
    for token in (" ", "\t", "\n", "\r", "　", ":", "："):
        text = text.replace(token, "")
    return text.replace("／", "/").replace("（", "(").replace("）", ")")


ALIAS_TO_CANONICAL = {
    _normalize_header(alias): canonical
    for canonical, aliases in COLUMN_ALIASES.items()
    for alias in aliases
}


def _canonical_header(value: Any) -> str:
    normalized = _normalize_header(value)
    return ALIAS_TO_CANONICAL.get(normalized, normalized)


def _canonicalize_headers(headers: list[Any]) -> list[HeaderInfo]:
    result: list[HeaderInfo] = []
    seen_names: dict[str, int] = {}
    seen_standard: set[str] = set()
    for header in headers:
        raw_name = _clean(header)
        if not raw_name:
            raw_name = f"未命名列{len(result) + 1}"
        canonical = _canonical_header(header)
        display_name = raw_name
        seen_names[display_name] = seen_names.get(display_name, 0) + 1
        if seen_names[display_name] > 1:
            display_name = f"{display_name}({seen_names[display_name]})"
        standard_key = canonical if canonical in COLUMN_MAP and canonical not in seen_standard else None
        if standard_key:
            seen_standard.add(standard_key)
        result.append(HeaderInfo(name=display_name, canonical=canonical, standard_key=standard_key))
    return result


DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%Y%m%d%H%M%S",
    "%Y%m%d%H%M",
    "%Y%m%d",
)


def _parse_excel_serial(value: int | float) -> datetime | None:
    if not 1 <= float(value) <= 100000:
        return None
    try:
        parsed = from_excel(value)
    except Exception:
        return None
    if isinstance(parsed, datetime):
        return parsed
    if isinstance(parsed, date):
        return datetime(parsed.year, parsed.month, parsed.day)
    return None


def _normalize_date_text(value: str) -> str:
    text = value.strip()
    text = re.sub(r"\s+", " ", text)
    text = text.replace("上午", " ").replace("下午", " ")
    text = text.replace("年", "-").replace("月", "-").replace("日", " ")
    text = text.replace(".", "-").replace("/", "-").replace("T", " ")
    text = text.replace("点", ":").replace("时", ":").replace("分", ":").replace("秒", "")
    text = re.sub(r"星期[一二三四五六日天]", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r":\s*$", "", text)
    text = re.sub(r"([+-]\d{2}:?\d{2}|Z)$", "", text).strip()
    return text


def _date_meridiem(value: str) -> str:
    if any(term in value for term in ("下午", "晚上", "晚间", "傍晚")):
        return "pm"
    if any(term in value for term in ("上午", "凌晨", "早上", "早晨")):
        return "am"
    return ""


def _apply_meridiem(hour: int, meridiem: str) -> int:
    if meridiem == "pm" and hour < 12:
        return hour + 12
    if meridiem == "am" and hour == 12:
        return 0
    return hour


def _parse_dt(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        as_int = int(value)
        if float(value).is_integer() and 10000101 <= as_int <= 99991231:
            value = str(as_int)
        else:
            return _parse_excel_serial(value)
    raw_text = str(value)
    meridiem = _date_meridiem(raw_text)
    text = _normalize_date_text(raw_text)
    if not text:
        return None
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:\s+(\d{1,2})(?::(\d{1,2}))?(?::(\d{1,2}))?)?", text)
    if match:
        try:
            year, month, day = (int(match.group(index)) for index in (1, 2, 3))
            hour = _apply_meridiem(int(match.group(4) or 0), meridiem)
            minute = int(match.group(5) or 0)
            second = int(match.group(6) or 0)
            return datetime(year, month, day, hour, minute, second)
        except ValueError:
            return None
    compact = re.sub(r"\D", "", text)
    if len(compact) in {8, 12, 14}:
        text = compact
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        pass
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _is_empty_row(values: tuple[Any, ...] | list[Any]) -> bool:
    return all(_clean(value) == "" for value in values)


def _find_header_index(rows: list[tuple[Any, ...] | list[Any]]) -> int:
    first_non_empty = 0
    best_index = -1
    best_score = 0
    for index, row in enumerate(rows[:20]):
        if _is_empty_row(row):
            continue
        if first_non_empty == 0 and index != 0:
            first_non_empty = index
        values = [_canonical_header(item) for item in row if item is not None and _clean(item)]
        score = sum(1 for value in values if value in COLUMN_MAP)
        if score > best_score:
            best_index = index
            best_score = score
    if best_index >= 0:
        return best_index
    for index, row in enumerate(rows[:20]):
        if not _is_empty_row(row):
            return index
    raise ValueError("文件为空")


def _extract_headers_and_rows(raw_rows: list[tuple[Any, ...] | list[Any]]) -> tuple[list[HeaderInfo], list[tuple[Any, ...] | list[Any]]]:
    if not raw_rows:
        raise ValueError("文件为空")

    header_index = _find_header_index(raw_rows)
    headers = _canonicalize_headers(list(raw_rows[header_index]))
    return headers, [row for row in raw_rows[header_index + 1 :] if not _is_empty_row(row)]


def _read_csv_rows(path: Path) -> list[list[str]]:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as fh:
                return list(csv.reader(fh))
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV 编码无法识别，请另存为 UTF-8 或 Excel 文件后再导入")


def _read_rows(path: Path) -> tuple[list[HeaderInfo], list[tuple[Any, ...] | list[Any]]]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("仅支持 .xlsx、.xlsm、.xls、.csv 文件")

    if suffix == ".csv":
        return _extract_headers_and_rows(_read_csv_rows(path))
    else:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                raw_rows = list(ws.iter_rows(values_only=True))
                try:
                    return _extract_headers_and_rows(raw_rows)
                except ValueError:
                    continue
            raise ValueError("文件为空")
        except (zipfile.BadZipFile, OSError):
            raw_rows = _read_legacy_excel_rows(path)
    return _extract_headers_and_rows(raw_rows)


def _read_legacy_excel_rows(path: Path) -> list[list[Any]]:
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError("该文件是旧版 Excel/WPS 格式，请安装 xlrd 后再导入") from exc
    book = xlrd.open_workbook(str(path))
    for sheet_index in range(book.nsheets):
        sheet = book.sheet_by_index(sheet_index)
        rows: list[list[Any]] = []
        for row_index in range(sheet.nrows):
            values: list[Any] = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        values.append(datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode)))
                    except Exception:
                        values.append(cell.value)
                else:
                    values.append(cell.value)
            rows.append(values)
        try:
            _extract_headers_and_rows(rows)
            return rows
        except ValueError:
            continue
    raise ValueError("文件为空")


def _build_payloads(headers: list[HeaderInfo], rows: list[tuple[Any, ...] | list[Any]]) -> tuple[list[dict[str, Any]], int]:
    payloads: list[dict[str, Any]] = []
    failed = 0
    for values in rows:
        payload = {target: "" for target in COLUMN_MAP.values()}
        payload["closed_at"] = None
        payload["received_at"] = None
        extra_fields: list[dict[str, str]] = []
        row_values = list(values)
        row_headers = list(headers)
        if len(row_values) > len(row_headers):
            for index in range(len(row_headers), len(row_values)):
                row_headers.append(HeaderInfo(name=f"未命名列{index + 1}", canonical=f"未命名列{index + 1}"))
        for index, header in enumerate(row_headers):
            value = row_values[index] if index < len(row_values) else ""
            if header.standard_key:
                target = COLUMN_MAP[header.standard_key]
                payload[target] = _parse_dt(value) if target in {"closed_at", "received_at"} else _clean(value)
            else:
                extra_fields.append({"name": header.name, "value": _clean(value)})
        if not payload.get("title") and payload.get("content"):
            payload["title"] = payload["content"][:80]
        if not payload.get("content") and payload.get("title"):
            payload["content"] = payload["title"]
        payload["extra_fields"] = json.dumps(extra_fields, ensure_ascii=False)
        payloads.append(payload)
    return payloads, failed


def import_workbook(db: Session, file_path: str | Path, imported_by: str = "system", original_filename: str | None = None) -> ImportBatch:
    path = Path(file_path)
    batch = ImportBatch(filename=original_filename or path.name, source="excel", status="running", imported_by=imported_by)
    db.add(batch)
    db.commit()
    db.refresh(batch)

    try:
        headers, rows = _read_rows(path)
        batch.total_rows = len(rows)
        payloads, failed = _build_payloads(headers, rows)
        success = len(payloads)
        if not payloads:
            batch.success_rows = 0
            batch.failed_rows = failed
            batch.status = "failed"
            batch.error_message = "未导入任何有效数据，请检查表头下方是否存在数据行"
            db.commit()
            return batch

        for payload in payloads:
            db.add(WorkOrder(batch_id=batch.id, **payload))
        batch.success_rows = success
        batch.failed_rows = failed
        batch.status = "success"
        batch.error_message = ""
        db.commit()
    except Exception as exc:
        db.rollback()
        batch = db.get(ImportBatch, batch.id)
        batch.status = "failed"
        batch.error_message = str(exc)
        db.commit()
    return batch
