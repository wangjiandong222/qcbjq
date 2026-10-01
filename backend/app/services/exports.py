from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from sqlalchemy.orm import Session

from app.core.config import PROJECT_DIR
from app.models.entities import CaseClassification, Clue, ClusterMember, EventCluster, ExportTask, WorkOrder


EXPORT_DIR = PROJECT_DIR / "data" / "exports"
WAGE_EMPTY = "无"
WAGE_FIELD_ORDER = (
    "工程项目名称",
    "工程地点",
    "开工时间",
    "工人人数",
    "欠薪主体",
    "欠薪数额",
    "是否签订劳动合同",
)
WAGE_TERMS = ("工资", "欠薪", "工钱", "劳务费", "工程款", "薪资", "欠发", "拖欠", "未发", "讨薪")
MONEY_EXCLUDE_TERMS = ("罚款", "处罚", "罚金", "收费", "缴费", "房租", "押金", "退款", "赔偿", "医疗费", "水费", "电费", "物业费")
PROJECT_REJECT_TERMS = ("拖欠", "欠薪", "工资", "工钱", "劳务费", "反映", "投诉", "要求", "希望", "处理", "解决", "咨询", "农民工", "工人", "包工头")
SUBJECT_REJECT_TERMS = ("拖欠", "欠薪", "工资", "工钱", "劳务费", "投诉", "反映", "要求", "希望", "处理", "解决")
ALERT_MODULE_PRIORITY = {
    "public_interest": 0,
    "vulnerable": 1,
    "administrative": 2,
}


def _extra_fields(order: WorkOrder) -> list[dict[str, str]]:
    try:
        fields = json.loads(order.extra_fields or "[]")
    except json.JSONDecodeError:
        return []
    return [item for item in fields if isinstance(item, dict) and item.get("name")]


def order_export_row(order: WorkOrder) -> dict:
    row = {
        "工单编号": order.order_no,
        "工单类型": order.order_type,
        "问题分类": order.problem_category,
        "标签": order.tags,
        "标题": order.title,
        "主要内容": order.content,
        "工单状态": order.status,
        "来电人": order.caller_name,
        "来电人电话/账号": order.caller_phone,
        "被反映区": order.district,
        "被反映街乡镇": order.town,
        "办理结果": order.handling_result,
        "回复内容": order.reply_content,
        "处理受理方式": order.handling_method,
        "来电/受理时间": order.received_at.isoformat() if order.received_at else "",
        "办结时间": order.closed_at.isoformat() if order.closed_at else "",
        "主办单位": order.host_unit,
        "企业名称": order.company_name,
        "是否解决": order.is_resolved,
        "是否满意": order.satisfaction,
        "工单性质": order.order_nature,
        "村/社区": order.community,
        "小区点位": order.location_point,
    }
    for field in _extra_fields(order):
        name = str(field.get("name") or "").strip()
        if name and name not in row:
            row[name] = field.get("value", "")
    return row


def _extra_text(order: WorkOrder) -> str:
    return "；".join(f"{item.get('name')}：{item.get('value', '')}" for item in _extra_fields(order))


def _wage_source_text(order: WorkOrder) -> str:
    return "；".join(
        str(value or "")
        for value in [
            order.title,
            order.content,
            order.handling_result,
            order.reply_content,
            order.host_unit,
            order.company_name,
            order.location_point,
            _extra_text(order),
        ]
        if value
    )


def _truncate_at_terms(value: str, terms: tuple[str, ...]) -> str:
    indexes = [value.find(term) for term in terms if term in value and value.find(term) > 0]
    if indexes:
        return value[: min(indexes)]
    return value


def _clean_candidate(
    value: str | None,
    *,
    max_len: int = 50,
    reject_terms: tuple[str, ...] = (),
    strip_terms: tuple[str, ...] = (),
) -> str:
    if not value:
        return WAGE_EMPTY
    text = re.sub(r"\s+", "", str(value))
    text = re.split(r"[，,。；;\n\r]", text, maxsplit=1)[0]
    text = _truncate_at_terms(text, strip_terms)
    text = re.sub(r"^(?:为|是|在|位于|地址|地点|项目名称|工程名称|工地名称|施工地点|工程地点|欠薪主体|拖欠主体|欠薪单位|用工单位|施工单位|劳务公司|承包方|发包方|包工头|老板)[:：为是]*", "", text)
    text = text.strip(" ：:，,。；;、（）()[]【】\"'“”‘’")
    if not text or text in {"无", "不详", "未知", "未提供"}:
        return WAGE_EMPTY
    if len(text) > max_len:
        return WAGE_EMPTY
    if any(term in text for term in reject_terms):
        return WAGE_EMPTY
    return text


def _first_clean_match(
    patterns: tuple[str, ...],
    text: str,
    *,
    max_len: int = 50,
    reject_terms: tuple[str, ...] = (),
    strip_terms: tuple[str, ...] = (),
) -> str:
    for pattern in patterns:
        for match in re.finditer(pattern, text):
            value = next((group for group in match.groups() if group), match.group(0))
            candidate = _clean_candidate(value, max_len=max_len, reject_terms=reject_terms, strip_terms=strip_terms)
            if candidate != WAGE_EMPTY:
                return candidate
    return WAGE_EMPTY


def _extract_project_name(text: str) -> str:
    candidate = _first_clean_match(
        (
            r"(?:项目名称|工程名称|工地名称|项目名|工程名)[为是：:]*([^，。；;\n]{2,40})",
            r"(?:在|于)[^，。；;\n]{0,18}(?:一个叫|名为|叫)([^，。；;\n]{2,28}(?:住宅项目|项目|工程|工地|标段))(?:施工|务工|干活|工作|上班)?",
            r"(?:在|于)([^，。；;\n]{2,40}(?:项目|工程|工地|标段))(?:施工|务工|干活|工作|上班)",
            r"([^，。；;\n]{2,40}(?:项目|工程|工地|标段))(?:拖欠|欠发|未发|拒付|拖欠工资|欠薪).{0,8}(?:工资|工钱|劳务费|欠薪)",
        ),
        text,
        max_len=40,
        reject_terms=PROJECT_REJECT_TERMS,
    )
    if candidate == WAGE_EMPTY:
        return candidate
    if any(generic in candidate for generic in ("某项目", "某工程", "某工地", "某标段", "该项目", "该工程", "该工地", "本项目")):
        return WAGE_EMPTY
    return candidate


def _extract_location(order: WorkOrder, text: str) -> str:
    point = _clean_candidate(order.location_point, max_len=60)
    if point != WAGE_EMPTY:
        return point
    return _first_clean_match(
        (
            r"(?:在|位于)([\u4e00-\u9fa5]{2,18}(?:区|镇|乡|街道|村|社区))(?:一个叫|名为|叫)",
            r"(?:施工地点|工程地点|项目地点|工地地址|工程地址|地址)[为是：:]*([^，。；;\n]{2,60})",
            r"(?:位于|在)([^，。；;\n]{2,60}(?:区|镇|街|路|村|社区|小区|号院|工地|项目))(?:施工|务工|干活|工作|上班|建设)",
        ),
        text,
        max_len=60,
        reject_terms=("工资", "欠薪", "工钱", "劳务费", "投诉", "反映", "要求", "希望"),
        strip_terms=("开工时间", "开工", "进场", "拖欠", "欠薪", "工资", "工钱", "劳务费"),
    )


def _extract_start_time(text: str) -> str:
    return _first_clean_match(
        (
            r"(?:开工时间|开工|进场|开始干活|开始上班|开始施工|从)[为是：:]*([0-9]{4}年[0-9]{1,2}月(?:[0-9]{1,2}日)?|[0-9]{4}[-/][0-9]{1,2}(?:[-/][0-9]{1,2})?|[0-9]{1,2}月[0-9]{1,2}日|去年[0-9]{1,2}月|今年[0-9]{1,2}月|[0-9]{4}年)",
        ),
        text,
        max_len=18,
    )


def _extract_worker_count(text: str) -> str:
    for pattern in (
        r"([0-9一二两三四五六七八九十百千万]+(?:余|多)?(?:名|个)?(?:工人|农民工|工友|务工人员))",
        r"([一二两三四五六七八九十百千万]+几(?:名|个)?(?:工人|农民工|工友|务工人员|人))",
        r"(?:共|共有|涉及|拖欠|欠了|班组)([0-9一二两三四五六七八九十百千万]+(?:余|多)?|[一二两三四五六七八九十百千万]+几)(?:名|个)?(?:工人|农民工|工友|务工人员|人)",
    ):
        match = re.search(pattern, text)
        if not match:
            continue
        value = match.group(1)
        if re.search(r"(?:工人|农民工|工友|务工人员|人)$", value):
            return value
        return f"{value}名"
    return WAGE_EMPTY


def _extract_wage_subject(order: WorkOrder, text: str) -> str:
    if order.company_name and any(term in text for term in WAGE_TERMS):
        company = _clean_candidate(order.company_name, max_len=50, reject_terms=SUBJECT_REJECT_TERMS)
        if company != WAGE_EMPTY:
            return company
    return _first_clean_match(
        (
            r"((?:包工头|老板)[\u4e00-\u9fa5A-Za-z0-9（）()·\-]{1,20}?)(?:拖欠|欠发|未发|不给|拒付)",
            r"((?:包工头|老板)[\u4e00-\u9fa5A-Za-z0-9（）()·\-]{1,12}?)(?:从|自).{0,18}(?:工资|欠薪|工钱|未结清|没结清|未发)",
            r"(?:欠薪主体|拖欠主体|欠薪单位|拖欠工资的单位|用工单位|施工单位|劳务公司|承包方|发包方)[为是：:]*([^，。；;\n]{2,50})",
            r"([^，。；;\n]{2,50}(?:公司|劳务公司|集团|施工队|工程队|项目部))(?:拖欠|欠发|未发|拒付)",
        ),
        text,
        max_len=50,
        reject_terms=SUBJECT_REJECT_TERMS,
    )


def _extract_wage_amount(text: str) -> str:
    explicit_patterns = (
        r"(?:拖欠工资|欠工资|欠薪|欠工钱|拖欠工钱|拖欠劳务费|劳务费|工程款|未发工资|欠发工资)[^，。；;\n]{0,12}?([0-9]+(?:\.[0-9]+)?(?:余|多)?\s*(?:万元|万|元))",
        r"(?:拖欠工资|欠工资|欠薪|欠工钱|拖欠工钱|拖欠劳务费|劳务费|工程款|未发工资|欠发工资)[^，。；;\n]{0,12}?([0-9]+(?:\.[0-9]+)?[-至到][0-9]+(?:\.[0-9]+)?\s*(?:万元|万|元))",
        r"([0-9]+(?:\.[0-9]+)?(?:余|多)?\s*(?:万元|万|元))[^，。；;\n]{0,8}?(?:工资|欠薪|工钱|劳务费|工程款)",
        r"(?:拖欠工资|欠工资|欠薪|欠工钱|拖欠工钱|拖欠劳务费|劳务费|工程款|未发工资|欠发工资)[^，。；;\n]{0,12}?([一二两三四五六七八九十百千万]+(?:余|多)?(?:万元|万|元))",
        r"(?:拖欠工资|欠工资|欠薪|欠工钱|拖欠工钱|拖欠劳务费|劳务费|工程款|未发工资|欠发工资)[^，。；;\n]{0,12}?([一二两三四五六七八九十百千万]+[一二两三四五六七八九十]?万)",
    )
    for pattern in explicit_patterns:
        candidate = _first_clean_match((pattern,), text, max_len=18)
        if candidate != WAGE_EMPTY:
            return candidate

    money_pattern = r"([0-9]+(?:\.[0-9]+)?(?:余|多)?\s*(?:万元|万|元)|[一二两三四五六七八九十百千万]+(?:余|多)?(?:万元|万|元))"
    for match in re.finditer(money_pattern, text):
        start, end = match.span()
        before = text[max(0, start - 14) : start]
        after = text[end : min(len(text), end + 8)]
        window = text[max(0, start - 18) : min(len(text), end + 12)]
        has_wage_context = any(term in before or term in after for term in WAGE_TERMS)
        has_excluded_context = any(term in window for term in MONEY_EXCLUDE_TERMS)
        if has_wage_context and not has_excluded_context:
            return match.group(1).replace(" ", "")
    return WAGE_EMPTY


def _extract_contract_status(text: str) -> str:
    negative_terms = ("未签劳动合同", "未签订劳动合同", "没有劳动合同", "没有签订劳动合同", "没签劳动合同", "没签合同", "无劳动合同", "未签订合同")
    positive_terms = ("已签订劳动合同", "签订了劳动合同", "签了劳动合同", "签订劳动合同", "有劳动合同", "签了合同")
    if any(term in text for term in negative_terms):
        return "否"
    if any(term in text for term in positive_terms):
        return "是"
    return WAGE_EMPTY


def extract_wage_fields(order: WorkOrder) -> dict[str, str]:
    text = _wage_source_text(order)
    fields = {
        "工程项目名称": _extract_project_name(text),
        "工程地点": _extract_location(order, text),
        "开工时间": _extract_start_time(text),
        "工人人数": _extract_worker_count(text),
        "欠薪主体": _extract_wage_subject(order, text),
        "欠薪数额": _extract_wage_amount(text),
        "是否签订劳动合同": _extract_contract_status(text),
    }
    return {name: fields.get(name) or WAGE_EMPTY for name in WAGE_FIELD_ORDER}


def format_wage_core(fields: dict[str, str]) -> str:
    return "\n".join(f"{name}：{fields.get(name) or WAGE_EMPTY}" for name in WAGE_FIELD_ORDER)


def extract_wage_core(order: WorkOrder) -> str:
    return format_wage_core(extract_wage_fields(order))


def _classification_row(db: Session, order: WorkOrder, module: str | None) -> dict:
    row = order_export_row(order)
    if module:
        classification = (
            db.query(CaseClassification)
            .filter(CaseClassification.work_order_id == order.id, CaseClassification.module == module)
            .first()
        )
        if classification:
            row.update(
                {
                    "业务板块": module,
                    "分类": classification.category,
                    "风险等级": classification.priority,
                    "确认状态": classification.review_status,
                    "复核人": classification.reviewer,
                    "复核时间": classification.reviewed_at.isoformat() if classification.reviewed_at else "",
                    "复核备注": classification.review_note,
                    "规则命中": classification.rule_hits,
                    "证据说明": classification.evidence,
                }
            )
    return row


def _autosize_sheet(ws, wrap_headers: set[str] | None = None) -> None:
    wrap_headers = wrap_headers or set()
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 46)
        if column_cells[0].value in wrap_headers:
            ws.column_dimensions[column_cells[0].column_letter].width = 38
            for cell in column_cells[1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[cell.row].height = max(ws.row_dimensions[cell.row].height or 15, 118)


def _write_rows_to_sheet(ws, rows: list[dict], empty_message: str = "暂无符合条件的数据", wrap_headers: set[str] | None = None) -> None:
    headers = list(rows[0].keys()) if rows else ["提示"]
    ws.append(headers)
    if rows:
        for row in rows:
            ws.append([row.get(header, "") for header in headers])
    else:
        ws.append([empty_message])
    _autosize_sheet(ws, wrap_headers)


def _write_workbook(rows: list[dict], file_path: Path, wrap_headers: set[str] | None = None) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "导出结果"
    _write_rows_to_sheet(ws, rows, wrap_headers=wrap_headers)
    wb.save(file_path)
    return wb


def _mark_exported(orders: list[WorkOrder], username: str, context: str) -> None:
    now = datetime.utcnow()
    for order in orders:
        order.exported_at = now
        order.exported_by = username
        order.export_context = context


def _unique_ids(values: list[int] | None) -> list[int]:
    return [value for value in dict.fromkeys(values or []) if value]


def _category_order_ids(db: Session, module: str | None, category: str | None, hide_exported: bool) -> list[int]:
    if not module:
        return []
    query = db.query(WorkOrder.id).join(CaseClassification, CaseClassification.work_order_id == WorkOrder.id).filter(CaseClassification.module == module)
    if category:
        query = query.filter(CaseClassification.category == category)
    if hide_exported:
        query = query.filter(WorkOrder.exported_at.is_(None))
    return _unique_ids([row[0] for row in query.order_by(WorkOrder.id.asc()).all()])


def _cluster_by_order_id(db: Session, order_ids: list[int]) -> dict[int, EventCluster]:
    if not order_ids:
        return {}
    rows = (
        db.query(ClusterMember, EventCluster)
        .join(EventCluster, ClusterMember.cluster_id == EventCluster.id)
        .filter(ClusterMember.work_order_id.in_(order_ids))
        .all()
    )
    return {member.work_order_id: cluster for member, cluster in rows}


def _member_ids_for_clusters(db: Session, cluster_ids: list[int]) -> list[int]:
    if not cluster_ids:
        return []
    return _unique_ids(
        [
            row[0]
            for row in (
                db.query(ClusterMember.work_order_id)
                .filter(ClusterMember.cluster_id.in_(cluster_ids))
                .order_by(ClusterMember.cluster_id.asc(), ClusterMember.work_order_id.asc())
                .all()
            )
        ]
    )


def _representative_ids_for_clusters(db: Session, cluster_ids: list[int]) -> list[int]:
    if not cluster_ids:
        return []
    return _unique_ids(
        [
            row[0]
            for row in (
                db.query(EventCluster.representative_order_id)
                .filter(EventCluster.id.in_(cluster_ids), EventCluster.representative_order_id.is_not(None))
                .order_by(EventCluster.id.asc())
                .all()
            )
        ]
    )


def _orders_by_ids(db: Session, order_ids: list[int]) -> list[WorkOrder]:
    order_ids = _unique_ids(order_ids)
    if not order_ids:
        return []
    rows = db.query(WorkOrder).filter(WorkOrder.id.in_(order_ids)).all()
    by_id = {order.id: order for order in rows}
    return [by_id[order_id] for order_id in order_ids if order_id in by_id]


def _resolve_duplicate_export_ids(
    db: Session,
    work_order_ids: list[int],
    cluster_ids: list[int],
    representative_work_order_ids: list[int],
    duplicate_export_scope: str,
    identify_duplicates: bool,
) -> list[int]:
    base_ids = _unique_ids([*work_order_ids, *representative_work_order_ids])
    cluster_ids = _unique_ids(cluster_ids)
    if not identify_duplicates:
        if duplicate_export_scope == "all_members" and cluster_ids:
            return _unique_ids([*base_ids, *_member_ids_for_clusters(db, cluster_ids)])
        return base_ids

    cluster_by_order = _cluster_by_order_id(db, base_ids)
    all_cluster_ids = _unique_ids([*cluster_ids, *[cluster.id for cluster in cluster_by_order.values()]])
    clustered_order_ids = set(cluster_by_order.keys())
    normal_ids = [order_id for order_id in base_ids if order_id not in clustered_order_ids]
    if duplicate_export_scope == "all_members":
        cluster_order_ids = _member_ids_for_clusters(db, all_cluster_ids)
    else:
        cluster_order_ids = _representative_ids_for_clusters(db, all_cluster_ids)
    return _unique_ids([*normal_ids, *cluster_order_ids])


def _classification_export_key(classification: CaseClassification) -> tuple[int, float, float]:
    updated_at = classification.updated_at.timestamp() if classification.updated_at else 0
    return (
        ALERT_MODULE_PRIORITY.get(classification.module, 99),
        -float(classification.risk_score or 0),
        -updated_at,
    )


def _best_classification_rows(rows: list[tuple[CaseClassification, WorkOrder]]) -> list[tuple[CaseClassification, WorkOrder]]:
    grouped: dict[int, tuple[CaseClassification, WorkOrder]] = {}
    for classification, order in rows:
        existing = grouped.get(order.id)
        if not existing or _classification_export_key(classification) < _classification_export_key(existing[0]):
            grouped[order.id] = (classification, order)
    return sorted(grouped.values(), key=lambda pair: (pair[1].id, _classification_export_key(pair[0])))


def _duplicate_grouped_export_ids(
    db: Session,
    work_order_ids: list[int],
    cluster_ids: list[int],
    representative_work_order_ids: list[int],
) -> tuple[list[int], list[int]]:
    base_ids = _unique_ids([*work_order_ids, *representative_work_order_ids])
    cluster_ids = _unique_ids(cluster_ids)
    cluster_by_order = _cluster_by_order_id(db, base_ids)
    all_cluster_ids = _unique_ids([*cluster_ids, *[cluster.id for cluster in cluster_by_order.values()]])
    return base_ids, all_cluster_ids


def _write_duplicate_expanded_workbook(
    db: Session,
    *,
    task: ExportTask,
    export_type: str,
    username: str,
    work_order_ids: list[int],
    cluster_ids: list[int],
    representative_work_order_ids: list[int],
    module: str | None,
    category: str | None,
) -> ExportTask:
    base_ids, explicit_cluster_ids = _duplicate_grouped_export_ids(
        db,
        work_order_ids,
        cluster_ids,
        representative_work_order_ids,
    )
    cluster_by_base_order = _cluster_by_order_id(db, base_ids)
    explicit_clusters = (
        {
            cluster.id: cluster
            for cluster in db.query(EventCluster).filter(EventCluster.id.in_(explicit_cluster_ids)).order_by(EventCluster.id.asc()).all()
        }
        if explicit_cluster_ids
        else {}
    )
    emitted_order_ids: set[int] = set()
    emitted_cluster_ids: set[int] = set()
    rows: list[dict] = []
    orders: list[WorkOrder] = []

    def append_order(order: WorkOrder, cluster: EventCluster | None = None, sequence: int | str = "") -> None:
        if order.id in emitted_order_ids:
            return
        emitted_order_ids.add(order.id)
        orders.append(order)
        row = {
            "重复事件ID": cluster.id if cluster else "",
            "重复事件标题": cluster.title if cluster else "",
            "重复事件序号": sequence if cluster else "",
            "是否代表工单": "是" if cluster and cluster.representative_order_id == order.id else "否",
        }
        row.update(order_export_row(order))
        rows.append(row)

    def append_cluster(cluster: EventCluster) -> None:
        if cluster.id in emitted_cluster_ids:
            return
        emitted_cluster_ids.add(cluster.id)
        member_ids = _member_ids_for_clusters(db, [cluster.id])
        for index, order in enumerate(_orders_by_ids(db, member_ids), start=1):
            append_order(order, cluster, index)

    for order_id in base_ids:
        cluster = cluster_by_base_order.get(order_id)
        if cluster:
            append_cluster(cluster)
            continue
        order = db.get(WorkOrder, order_id)
        if order:
            append_order(order)
    for cluster in explicit_clusters.values():
        append_cluster(cluster)

    file_path = EXPORT_DIR / f"{export_type}_{task.id}.xlsx"
    _write_workbook(rows, file_path)
    if orders:
        _mark_exported(orders, username, export_type)
    task.status = "success"
    task.file_path = str(file_path)
    db.commit()
    db.refresh(task)
    return task


def create_selected_export(
    db: Session,
    export_type: str,
    username: str,
    work_order_ids: list[int] | None = None,
    cluster_ids: list[int] | None = None,
    representative_work_order_ids: list[int] | None = None,
    duplicate_export_scope: str = "representative",
    module: str | None = None,
    category: str | None = None,
    identify_duplicates: bool = False,
    hide_exported: bool = False,
) -> ExportTask:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    task = ExportTask(export_type=export_type, status="running", created_by=username)
    db.add(task)
    db.commit()
    db.refresh(task)

    rows: list[dict] = []
    orders: list[WorkOrder] = []
    work_order_ids = _unique_ids(work_order_ids)
    cluster_ids = _unique_ids(cluster_ids)
    representative_work_order_ids = _unique_ids(representative_work_order_ids)
    if not work_order_ids and not representative_work_order_ids and module and export_type not in {"clusters", "clusters_selected", "alerts_selected"}:
        work_order_ids = _category_order_ids(db, module, category, hide_exported)
    if identify_duplicates and export_type not in {"clusters", "clusters_selected", "alerts_selected", "wage"}:
        return _write_duplicate_expanded_workbook(
            db,
            task=task,
            export_type=export_type,
            username=username,
            work_order_ids=work_order_ids,
            cluster_ids=cluster_ids,
            representative_work_order_ids=representative_work_order_ids,
            module=module,
            category=category,
        )
    work_order_ids = _resolve_duplicate_export_ids(db, work_order_ids, cluster_ids, representative_work_order_ids, duplicate_export_scope, identify_duplicates)
    if export_type in {"clusters", "clusters_selected"}:
        cluster_query = db.query(EventCluster)
        if cluster_ids:
            cluster_query = cluster_query.filter(EventCluster.id.in_(cluster_ids))
        rows = [
            {
                "屡诉未决事件ID": item.id,
                "屡诉未决事件标题": item.title,
                "代表工单ID": item.representative_order_id or "",
                "点位": item.location_point,
                "街乡镇": item.town,
                "投诉次数": item.complaint_count,
                "未解决次数": item.unresolved_count,
                "不满意次数": item.dissatisfied_count,
                "风险等级": item.risk_level,
                "识别理由": item.match_reason,
            }
            for item in cluster_query.order_by(EventCluster.id.asc()).all()
        ]
        file_path = EXPORT_DIR / f"{export_type}_{task.id}.xlsx"
        wb = Workbook()
        summary = wb.active
        summary.title = "事件汇总"
        headers = list(rows[0].keys()) if rows else ["提示"]
        summary.append(headers)
        if rows:
            for row in rows:
                summary.append([row.get(header, "") for header in headers])
        else:
            summary.append(["暂无符合条件的数据"])
        members_ws = wb.create_sheet("成员投诉")
        member_rows: list[dict] = []
        member_orders: list[WorkOrder] = []
        if cluster_ids:
            member_query = (
                db.query(EventCluster, ClusterMember, WorkOrder)
                .join(ClusterMember, ClusterMember.cluster_id == EventCluster.id)
                .join(WorkOrder, ClusterMember.work_order_id == WorkOrder.id)
                .filter(EventCluster.id.in_(cluster_ids))
                .order_by(EventCluster.id.asc(), WorkOrder.id.asc())
            )
        else:
            member_query = (
                db.query(EventCluster, ClusterMember, WorkOrder)
                .join(ClusterMember, ClusterMember.cluster_id == EventCluster.id)
                .join(WorkOrder, ClusterMember.work_order_id == WorkOrder.id)
                .order_by(EventCluster.id.asc(), WorkOrder.id.asc())
            )
        for cluster, _, order in member_query.all():
            row = {"屡诉未决事件ID": cluster.id, "屡诉未决事件标题": cluster.title, "是否代表工单": "是" if cluster.representative_order_id == order.id else "否"}
            row.update(order_export_row(order))
            member_rows.append(row)
            member_orders.append(order)
        member_headers = list(member_rows[0].keys()) if member_rows else ["提示"]
        members_ws.append(member_headers)
        if member_rows:
            for row in member_rows:
                members_ws.append([row.get(header, "") for header in member_headers])
        else:
            members_ws.append(["暂无成员投诉"])
        for ws in (summary, members_ws):
            for column_cells in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in column_cells)
                ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 46)
        if member_orders:
            _mark_exported(member_orders, username, export_type)
        wb.save(file_path)
        task.status = "success"
        task.file_path = str(file_path)
        db.commit()
        db.refresh(task)
        return task
    elif export_type == "alerts_selected":
        query = db.query(CaseClassification, WorkOrder).join(WorkOrder, CaseClassification.work_order_id == WorkOrder.id)
        if work_order_ids:
            query = query.filter(WorkOrder.id.in_(work_order_ids))
        all_classification_rows = query.all()
        high_risk_by_order = {
            order.id: (classification, order)
            for classification, order in _best_classification_rows([row for row in all_classification_rows if row[0].priority == "高"])
        }
        any_by_order = {
            order.id: (classification, order)
            for classification, order in _best_classification_rows(all_classification_rows)
        }
        target_orders = _orders_by_ids(db, work_order_ids) if work_order_ids else [order for _, order in _best_classification_rows(all_classification_rows)]
        cluster_by_order = _cluster_by_order_id(db, [order.id for order in target_orders]) if identify_duplicates else {}
        cluster_sequences: dict[tuple[int, int], int] = {}
        if cluster_by_order:
            for cluster_id in _unique_ids([cluster.id for cluster in cluster_by_order.values()]):
                for index, order_id in enumerate(_member_ids_for_clusters(db, [cluster_id]), start=1):
                    cluster_sequences[(cluster_id, order_id)] = index
        for order in target_orders:
            pair = high_risk_by_order.get(order.id) or any_by_order.get(order.id)
            classification = pair[0] if pair else None
            row = order_export_row(order)
            cluster = cluster_by_order.get(order.id)
            if cluster:
                row = {
                    "重复事件ID": cluster.id,
                    "重复事件标题": cluster.title,
                    "重复事件序号": cluster_sequences.get((cluster.id, order.id), ""),
                    "是否代表工单": "是" if cluster.representative_order_id == order.id else "否",
                    **row,
                }
            elif identify_duplicates:
                row = {
                    "重复事件ID": "",
                    "重复事件标题": "",
                    "重复事件序号": "",
                    "是否代表工单": "否",
                    **row,
                }
            row.update(
                {
                    "业务板块": classification.module if classification else "",
                    "分类": classification.category if classification else "",
                    "风险等级": classification.priority if classification else "",
                    "确认状态": classification.review_status if classification else "",
                    "规则命中": classification.rule_hits if classification else "",
                    "证据说明": classification.evidence if classification else "",
                    "复核人": classification.reviewer if classification else "",
                    "复核时间": classification.reviewed_at.isoformat() if classification and classification.reviewed_at else "",
                    "复核备注": classification.review_note if classification else "",
                }
            )
            rows.append(row)
            orders.append(order)
    else:
        query = db.query(WorkOrder)
        if work_order_ids:
            query = query.filter(WorkOrder.id.in_(work_order_ids))
        elif module:
            query = query.join(CaseClassification, CaseClassification.work_order_id == WorkOrder.id).filter(CaseClassification.module == module)
            if category:
                query = query.filter(CaseClassification.category == category)
            if hide_exported:
                query = query.filter(WorkOrder.exported_at.is_(None))
        orders = query.order_by(WorkOrder.id.asc()).all()
        for order in orders:
            if export_type == "wage":
                original = order_export_row(order)
                row = {"工单编号": order.order_no, "核心内容提炼": extract_wage_core(order)}
                for key, value in original.items():
                    if key != "工单编号":
                        row[key] = value
                rows.append(row)
            else:
                rows.append(_classification_row(db, order, module))

    file_path = EXPORT_DIR / f"{export_type}_{task.id}.xlsx"
    _write_workbook(rows, file_path, wrap_headers={"核心内容提炼"} if export_type == "wage" else set())
    if orders:
        _mark_exported(orders, username, export_type)
    task.status = "success"
    task.file_path = str(file_path)
    db.commit()
    db.refresh(task)
    return task


def create_export(db: Session, export_type: str, username: str) -> ExportTask:
    if export_type == "confirmed_domains":
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        task = ExportTask(export_type=export_type, status="running", created_by=username)
        db.add(task)
        db.commit()
        db.refresh(task)
        rows: list[dict] = []
        orders: list[WorkOrder] = []
        for classification in db.query(CaseClassification).filter(CaseClassification.review_status == "已确认").all():
            if not classification.work_order:
                continue
            row = order_export_row(classification.work_order)
            row.update(
                {
                    "人工确认成案领域": classification.category,
                    "风险等级": classification.priority,
                    "确认状态": classification.review_status,
                    "复核人": classification.reviewer,
                    "复核时间": classification.reviewed_at.isoformat() if classification.reviewed_at else "",
                    "复核备注": classification.review_note,
                    "规则命中": classification.rule_hits,
                    "证据说明": classification.evidence,
                }
            )
            rows.append(row)
            orders.append(classification.work_order)
        if not rows:
            for clue, order in db.query(Clue, WorkOrder).join(WorkOrder, Clue.work_order_id == WorkOrder.id).filter(Clue.status == "已确认").all():
                row = order_export_row(order)
                row.update(
                    {
                        "人工确认成案领域": clue.predicted_domain,
                        "风险等级": clue.priority,
                        "确认状态": clue.status,
                        "复核人": clue.reviewer,
                        "复核时间": clue.reviewed_at.isoformat() if clue.reviewed_at else "",
                        "复核备注": clue.review_note,
                        "规则命中": clue.rule_hits,
                        "证据说明": clue.evidence,
                    }
                )
                rows.append(row)
                orders.append(order)
        file_path = EXPORT_DIR / f"{export_type}_{task.id}.xlsx"
        _write_workbook(rows, file_path)
        if orders:
            _mark_exported(orders, username, export_type)
        task.status = "success"
        task.file_path = str(file_path)
        db.commit()
        db.refresh(task)
        return task
    return create_selected_export(db, export_type, username)
